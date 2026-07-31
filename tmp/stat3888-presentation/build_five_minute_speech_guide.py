"""Build the STAT3888 five-minute speech and principles guide."""

from pathlib import Path
import re
import sys
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from five_minute_speech_content import (
    PRINCIPLE_SECTIONS,
    PROJECT_FACTS,
    SPEECH_SECTIONS,
    VIVA_QA,
)


WORKTREE_ROOT = Path(__file__).resolve().parents[2]
DELIVERABLE_DIR = (
    WORKTREE_ROOT
    / "deliverables"
    / "STAT3888_Toronto_Bicycle_Theft_Presentation"
)
OUTPUT = DELIVERABLE_DIR / "STAT3888_五分钟演讲稿与原理解析_滚动交叉验证版.docx"
SLIDE_DIR = (
    WORKTREE_ROOT
    / "deliverables"
    / "STAT3888_Toronto_Bicycle_Theft_Presentation"
    / "STAT3888_Toronto_Bicycle_Theft_Cross_Validated_Embedded_Video"
)

DOCUMENT_RUNTIME_ROOT = Path(
    "/Users/liumingyuan/.codex/plugins/cache/openai-primary-runtime/documents"
)
DOCUMENT_RUNTIME_CANDIDATES = sorted(
    DOCUMENT_RUNTIME_ROOT.glob("*/skills/documents")
)
if not DOCUMENT_RUNTIME_CANDIDATES:
    raise RuntimeError("No bundled documents runtime is available.")
SKILL_DIR = DOCUMENT_RUNTIME_CANDIDATES[-1]
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 33, 41)
MUTED = RGBColor(87, 98, 110)
TEAL = RGBColor(0, 127, 130)
PURPLE = RGBColor(98, 36, 109)
RED = RGBColor(155, 28, 28)
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E8F3F4"
LIGHT_PURPLE = "F1EAF3"
LIGHT_GRAY = "F2F4F7"
WHITE = RGBColor(255, 255, 255)


def set_font(
    run,
    size=None,
    bold=None,
    italic=None,
    color=INK,
    ascii_font="Arial",
    east_asia="Hiragino Sans GB",
):
    text = run.text or ""
    font_name = east_asia if any("\u3400" <= char <= "\u9fff" for char in text) else ascii_font
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), font_name)
    r_pr.rFonts.set(qn("w:hAnsi"), font_name)
    r_pr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_font(run, size=9, color=MUTED)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in [
        ("Heading 1", 17, BLUE, 16, 8),
        ("Heading 2", 13.5, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ]:
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_paragraph.add_run("STAT3888 · Toronto Bicycle Theft · Task 1 + Task 4")
    set_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_paragraph.add_run("五分钟演讲稿与原理解析  ·  ")
    set_font(run, size=8.5, color=MUTED)
    add_page_number(footer_paragraph)


def normalise_office_theme_fonts(docx_path):
    """Replace Microsoft-only theme defaults with fonts installed on this Mac."""
    temporary_path = docx_path.with_suffix(".font-fix.docx")
    with ZipFile(docx_path, "r") as source, ZipFile(
        temporary_path, "w", compression=ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/theme/theme1.xml":
                data = re.sub(
                    rb'typeface="[^"]*"',
                    b'typeface="Arial"',
                    data,
                )
            elif item.filename in {"word/settings.xml", "word/fontTable.xml"}:
                data = data.replace(
                    b'w:eastAsia="ja-JP"',
                    b'w:eastAsia="zh-CN"',
                ).replace(
                    b"Cambria Math",
                    b"Times New Roman",
                ).replace(
                    b"Calibri",
                    b"Arial",
                ).replace(
                    b"Cambria",
                    b"Times New Roman",
                ).replace(
                    "ＭＳ 明朝".encode("utf-8"),
                    b"Arial",
                ).replace(
                    "ＭＳ ゴシック".encode("utf-8"),
                    b"Arial",
                )
            target.writestr(item, data)
    temporary_path.replace(docx_path)


def add_para(
    document,
    text="",
    size=10.5,
    bold=False,
    italic=False,
    color=INK,
    before=0,
    after=5,
    line=1.22,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    keep=False,
):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_with_next = keep
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def add_callout(document, label, text, fill=LIGHT_TEAL, accent=TEAL, after=8):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.14)
    paragraph.paragraph_format.right_indent = Inches(0.14)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.18
    set_paragraph_shading(paragraph, fill)
    label_run = paragraph.add_run(f"{label}\n")
    set_font(label_run, size=9, bold=True, color=accent)
    text_run = paragraph.add_run(text)
    set_font(text_run, size=10.5, bold=label in {"本页核心句", "评分安全线"})
    return paragraph


def add_labelled_para(document, label, text, label_color=DARK_BLUE, after=5):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.22
    label_run = paragraph.add_run(label)
    set_font(label_run, size=10.5, bold=True, color=label_color)
    text_run = paragraph.add_run(text)
    set_font(text_run, size=10.5)
    return paragraph


def add_table(document, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, header_fill)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_font(run, size=9, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            if index > 0 and len(str(text)) <= 20:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(text))
            set_font(run, size=8.8)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=9605,
        indent_dxa=120,
        cell_margins_dxa={"top": 90, "bottom": 90, "start": 110, "end": 110},
    )
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_slide_image(document, slide_number):
    image_path = SLIDE_DIR / f"slide-{slide_number}.png"
    if not image_path.exists():
        add_callout(
            document,
            "PPT缩略图",
            f"未找到第{slide_number}页渲染图；正式内容仍可独立使用。",
            fill=LIGHT_GRAY,
            accent=MUTED,
        )
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(6.0))
    inline._inline.docPr.set("descr", f"STAT3888 presentation slide {slide_number}")


def add_cover(document):
    add_para(
        document,
        "STAT3888 · SPATIOTEMPORAL DATA SCIENCE",
        size=10.5,
        bold=True,
        color=TEAL,
        before=56,
        after=20,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        document,
        "Toronto Bicycle Theft",
        size=29,
        bold=True,
        color=RGBColor(32, 55, 72),
        after=5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        document,
        "五分钟演讲稿与方法原理解析",
        size=17,
        color=RGBColor(43, 81, 99),
        after=25,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_callout(
        document,
        "使用方式",
        "Part I 可以直接照着练习；Part II 用于真正理解基函数回归；Part III 用于答辩前快速复习。",
        fill=LIGHT_TEAL,
        accent=TEAL,
        after=15,
    )
    add_labelled_para(document, "姓名：", "____________________________")
    add_labelled_para(document, "学号：", "____________________________")
    add_labelled_para(document, "项目范围：", "Task 1 时空可视化 + Task 4 基函数线性回归")
    add_labelled_para(document, "正式时长：", "5分钟（含指图、停顿和约5秒视频展示）")
    add_labelled_para(
        document,
        "数据：",
        "Toronto bicycle theft reports，2014–2023，31,833条事件，140个社区",
    )
    add_callout(
        document,
        "评分安全线",
        "Basis OLS 是Task 4标准模型；Basis Ridge是稳定化扩展。R²不是准确率，空间相关不是因果。",
        fill=LIGHT_PURPLE,
        accent=PURPLE,
        after=10,
    )
    document.add_page_break()


def add_quick_map(document):
    document.add_heading("使用说明与五分钟路线图", level=1)
    add_para(
        document,
        "先完整朗读两遍。第三遍打开PPT，只在“指图动作”规定的位置指图，不要逐字朗读坐标轴或图题。",
    )
    rows = []
    for section in SPEECH_SECTIONS:
        start = f"{section['start_seconds']//60}:{section['start_seconds']%60:02d}"
        end = f"{section['end_seconds']//60}:{section['end_seconds']%60:02d}"
        rows.append(
            (
                f"第{section['slide']}页",
                f"{start}–{end}",
                section["title"],
                section["core_sentence"],
            )
        )
    add_table(
        document,
        ["PPT", "时间", "任务", "必须让听众记住"],
        rows,
        [850, 900, 2050, 5805],
    )
    add_callout(
        document,
        "计时建议",
        "英文正式稿共697词。以每分钟约140词的清晰学术语速，并把视频展示控制在约5秒，可在约5分钟内完成。",
        fill=LIGHT_GRAY,
        accent=DARK_BLUE,
    )
    document.add_heading("正式演讲必须保持的四个区分", level=2)
    for text in [
        "Task 1负责发现和展示模式；Task 4负责建立预测模型。",
        "Basis OLS满足任务要求；Basis Ridge是在同一线性模型上的L2正则化扩展。",
        "OLS测试RMSE最低；Ridge的MAE更低且强调系数稳定性，不能说Ridge在所有指标上最好。",
        "reported theft counts是报告数量，不等于真实风险，也不证明因果关系。",
    ]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(text)
        set_font(run, size=10.5)
    document.add_page_break()


def add_speech_part(document):
    document.add_heading("Part I · 五分钟逐页演讲稿（English）", level=1)
    add_para(
        document,
        "下列内容按最终四页PPT排列。方括号中的动作不需要说出口。",
        color=MUTED,
        italic=True,
        after=9,
    )
    for index, section in enumerate(SPEECH_SECTIONS):
        start = f"{section['start_seconds']//60}:{section['start_seconds']%60:02d}"
        end = f"{section['end_seconds']//60}:{section['end_seconds']%60:02d}"
        heading = document.add_heading(
            f"第{section['slide']}页 · {section['title']}  |  {start}–{end}",
            level=2,
        )
        if index:
            heading.paragraph_format.page_break_before = True
        add_slide_image(document, section["slide"])
        add_callout(
            document,
            "指图动作",
            section["pointing_cues"],
            fill=LIGHT_GRAY,
            accent=DARK_BLUE,
            after=7,
        )
        add_para(
            document,
            section["script"],
            size=10.5 if section["slide"] == 3 else 11.2,
            line=1.20 if section["slide"] == 3 else 1.32,
            after=8,
        )
        add_callout(
            document,
            "本页核心句",
            section["core_sentence"],
            fill=LIGHT_TEAL,
            accent=TEAL,
            after=5,
        )


def add_principles_part(document):
    heading = document.add_heading("Part II · 方法原理解析", level=1)
    heading.paragraph_format.page_break_before = True
    add_para(
        document,
        "每个概念都按“直觉 → 数学表达 → 本项目如何实现 → 为什么适合”展开。"
        "先理解直觉，再看公式中的每个符号。",
        after=10,
    )
    for index, item in enumerate(PRINCIPLE_SECTIONS):
        heading = document.add_heading(item["title"], level=2)
        if index in {3, 6}:
            heading.paragraph_format.page_break_before = True
        add_labelled_para(document, "直觉：", item["intuition"])
        add_callout(
            document,
            "数学表达",
            item["formula"],
            fill=LIGHT_PURPLE,
            accent=PURPLE,
            after=5,
        )
        add_labelled_para(document, "本项目如何实现：", item["project_use"])
        add_labelled_para(document, "为什么适合：", item["why_suitable"], after=8)

    metric_heading = document.add_heading("模型结果应如何准确表述", level=2)
    metric_heading.paragraph_format.page_break_before = True
    metric_rows = [
        ("Global mean", "2.170", "4.164", "-0.006", "比测试集均值还差"),
        ("Neighbourhood mean", "1.391", "2.699", "0.577", "有空间差异但无时间变化"),
        ("Basis OLS", "1.120", "1.966", "0.776", "最低RMSE、最高R²"),
        ("Basis Ridge", "1.054", "1.976", "0.774", "最低MAE、系数更稳定"),
    ]
    add_table(
        document,
        ["模型", "MAE", "RMSE", "R²", "正确解释"],
        metric_rows,
        [1800, 900, 900, 900, 5105],
    )
    add_callout(
        document,
        "推荐说法",
        "Basis OLS在2023测试集取得最低RMSE 1.97和最高R² 0.776；Basis Ridge的RMSE为1.98、"
        "MAE更低，并通过滚动交叉验证选择λ，因此作为主要预测与残差展示模型。",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )
    document.add_heading("一分钟原理复述", level=2)
    recap = [
        "先把单条事件汇总为社区—月份计数，使时间和空间进入同一观察单位。",
        "用趋势图、热点图、热力图和动画分别回答什么时候高、哪里高、哪里在什么时候变高。",
        "用log(1+y)压缩极端计数，再在预测后返回原始计数尺度评价。",
        "用B-spline、sine/cosine和RBF把长期、季节和空间结构变成设计矩阵的列。",
        "Basis OLS完成Task 4；Basis Ridge在同一线性模型上加入L2惩罚以提高稳定性。",
        "用5个滚动时间折选择Ridge的λ，再用完全未参与调参的2023测试，并与简单基线比较。",
    ]
    for text in recap:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(text)
        set_font(run, size=10.2)


def add_viva_part(document):
    heading = document.add_heading("Part III · 答辩速查", level=1)
    heading.paragraph_format.page_break_before = True
    add_para(
        document,
        "先回答“20秒版本”。只有老师继续追问时，再使用“进一步解释”。",
        color=MUTED,
        italic=True,
        after=9,
    )
    for index, item in enumerate(VIVA_QA):
        heading = document.add_heading(item["question"], level=2)
        if index in {4, 7}:
            heading.paragraph_format.page_break_before = True
        add_callout(
            document,
            "20秒回答",
            item["short_answer"],
            fill=LIGHT_TEAL,
            accent=TEAL,
            after=5,
        )
        add_labelled_para(document, "进一步解释：", item["deeper_answer"], after=8)

    document.add_heading("上台前最后检查", level=2)
    checks = [
        "能在不看稿的情况下说出研究问题和两条核心发现。",
        "能解释热力图的行、列和颜色分别代表什么。",
        "能用“设计矩阵的列”解释基函数，而不是只背术语。",
        "能说出5个滚动验证年、最终训练期和2023测试集各自的目的。",
        "能准确区分OLS与Ridge的结果和选择理由。",
        "不会把R²说成预测准确率，也不会把空间热点说成因果。",
    ]
    for text in checks:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(text)
        set_font(run, size=10.5)


def build_document(output_path=OUTPUT):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)
    add_quick_map(document)
    add_speech_part(document)
    add_principles_part(document)
    add_viva_part(document)

    document.core_properties.title = "STAT3888 五分钟演讲稿与原理解析"
    document.core_properties.subject = "Task 1 spatiotemporal visualization and Task 4 basis-function regression"
    document.core_properties.author = "Name: ____________________"
    document.core_properties.keywords = "STAT3888, spatiotemporal, basis functions, OLS, Ridge"
    document.save(output_path)
    normalise_office_theme_fonts(output_path)
    return output_path


if __name__ == "__main__":
    result = build_document()
    print(result)
